"""Escritura del FGR-06 (CORPOBOYACÁ) – Parte A "Inventario Forestal al 100%".

Se parte de la plantilla limpia en /plantilla y se escribe una fila por árbol,
con las mismas celdas combinadas del formato oficial. Si algún árbol tiene más
de 3 fustes, se insertan columnas "DAP D (cm)", "DAP E (cm)"... después de DAP C.
"""

from __future__ import annotations

import io
import string
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from .calculos import decimal_a_gms
from .excel_utils import anchos_columnas, copiar_estilo, insertar_columnas, normalizar_bordes_combinadas
from .lector import Inventario

PLANTILLA = Path(__file__).resolve().parent.parent / "plantilla" / "FGR-06_v7_plantilla.xlsx"
HOJA_PARTE_A = 0          # índice de la hoja "Parte A. Inv. al 100%"
FILA_INICIO = 11          # primera fila de árboles
COL_DAP_C = 16            # columna P
FUSTES_BASE = 3           # A, B y C vienen en el formato

_delgada = Side(style="thin")
BORDE = Border(left=_delgada, right=_delgada, top=_delgada, bottom=_delgada)
FUENTE = Font(name="Arial", size=8)
FUENTE_TOTAL = Font(name="Arial", size=8, bold=True)
CENTRO = Alignment(horizontal="center", vertical="center")
IZQUIERDA = Alignment(horizontal="left", vertical="center")


def _grupos(extra: int) -> dict[str, tuple[int, int]]:
    """Columnas (inicio, fin) de cada campo, según los fustes extra insertados."""
    k = extra
    g = {
        "n": (1, 2),
        "nombre": (3, 11),
        "dap_1": (12, 14),
        "dap_2": (15, 15),
        "dap_3": (16, 16),
    }
    for i in range(extra):
        g[f"dap_{4 + i}"] = (17 + i, 17 + i)
    g.update({
        "ht": (17 + k, 18 + k),
        "vol": (19 + k, 21 + k),
        "lat_g": (22 + k, 24 + k),
        "lat_m": (25 + k, 27 + k),
        "lat_s": (28 + k, 30 + k),
        "lon_g": (31 + k, 33 + k),
        "lon_m": (34 + k, 36 + k),
        "lon_s": (37 + k, 39 + k),
    })
    g["id"] = (40 + k, 40 + k)
    return g


FORMATOS = {"dap": "0.00", "ht": "0", "vol": "0.000", "lat_s": "0.00", "lon_s": "0.00"}


def filas_fgr(inv: Inventario, decimales_seg: int = 2) -> tuple[list[dict], list[str]]:
    """Convierte los árboles a filas del FGR-06. Devuelve (filas, avisos)."""
    filas, avisos, coords = [], [], {}
    for a in inv.arboles:
        lat = decimal_a_gms(a.lat, decimales_seg)
        lon = decimal_a_gms(a.lon, decimales_seg)
        clave = (lat, lon)
        if clave in coords:
            avisos.append(
                f"ID {a.id} y ID {coords[clave]} quedan con la misma coordenada en grados-minutos-segundos."
            )
        coords.setdefault(clave, a.id)
        filas.append({
            "id": a.id,
            "nombre": a.nombre_comun,
            "cientifico": a.nombre_cientifico,
            "daps": a.daps_cm,
            "ht": a.altura,
            "vol": a.vt,
            "lat": lat,
            "lon": lon,
        })
    return filas, avisos


def _escribir(ws, fila, grupo, valor, formato=None, fuente=FUENTE, alineacion=CENTRO):
    c1, c2 = grupo
    for c in range(c1, c2 + 1):
        celda = ws.cell(fila, c)
        celda.border = BORDE
        celda.font = fuente
        celda.alignment = alineacion
    anclaje = ws.cell(fila, c1)
    anclaje.value = valor
    if formato:
        anclaje.number_format = formato
    if c2 > c1:
        ws.merge_cells(start_row=fila, start_column=c1, end_row=fila, end_column=c2)


def generar_fgr06(inv: Inventario, decimales_seg: int = 2, plantilla: Path | str = PLANTILLA) -> bytes:
    """Genera el FGR-06 y lo devuelve como bytes (.xlsx)."""
    wb = openpyxl.load_workbook(plantilla)
    ws = wb.worksheets[HOJA_PARTE_A]

    # --- Fustes adicionales -------------------------------------------------
    extra = max(0, inv.max_fustes - FUSTES_BASE)
    if extra:
        ancho_p = anchos_columnas(ws, COL_DAP_C)[COL_DAP_C]
        insertar_columnas(ws, COL_DAP_C + 1, extra, ancho=ancho_p)
        for i in range(extra):
            col = COL_DAP_C + 1 + i
            for r in (8, 9, 10):
                copiar_estilo(ws.cell(r, COL_DAP_C), ws.cell(r, col))
            ws.cell(8, col).value = f"DAP {string.ascii_uppercase[FUSTES_BASE + i]} (cm)"
            ws.merge_cells(start_row=8, start_column=col, end_row=10, end_column=col)

    g = _grupos(extra)
    # Encabezados DAP con la misma fuente que "DAP A"
    for j in range(2, FUSTES_BASE + extra + 1):
        ws.cell(8, g[f"dap_{j}"][0]).font = copy(ws.cell(8, g["dap_1"][0]).font)
    ultima_col = g["lon_s"][1]
    col_id = g["id"][0]

    # --- Filas de árboles ---------------------------------------------------
    filas, _ = filas_fgr(inv, decimales_seg)
    n_fustes = FUSTES_BASE + extra
    for i, f in enumerate(filas):
        r = FILA_INICIO + i
        ws.row_dimensions[r].height = 15
        _escribir(ws, r, g["n"], 1, "0")
        _escribir(ws, r, g["nombre"], f["nombre"])
        for j in range(n_fustes):
            dap = f["daps"][j] if j < len(f["daps"]) else 0
            _escribir(ws, r, g[f"dap_{j + 1}"], dap, FORMATOS["dap"])
        _escribir(ws, r, g["ht"], f["ht"], FORMATOS["ht"] if float(f["ht"]).is_integer() else "0.0")
        _escribir(ws, r, g["vol"], f["vol"], FORMATOS["vol"])
        for pre, (gr, mi, se) in (("lat", f["lat"]), ("lon", f["lon"])):
            _escribir(ws, r, g[f"{pre}_g"], gr, "0")
            _escribir(ws, r, g[f"{pre}_m"], mi, "0")
            _escribir(ws, r, g[f"{pre}_s"], se, FORMATOS[f"{pre}_s"])
        # ID oculto (columna agrupada; se ve con el botón "+")
        _escribir(ws, r, g["id"], f["id"])

    ultima_fila = FILA_INICIO + len(filas) - 1

    # --- Totales y nota -----------------------------------------------------
    r_tot = ultima_fila + 1
    ws.row_dimensions[r_tot].height = 15
    L = get_column_letter
    _escribir(ws, r_tot, g["n"], f"=SUM({L(1)}{FILA_INICIO}:{L(1)}{ultima_fila})", "0", FUENTE_TOTAL)
    _escribir(ws, r_tot, (g["nombre"][0], g["ht"][1]), "TOTAL", fuente=FUENTE_TOTAL)
    col_vol = L(g["vol"][0])
    _escribir(ws, r_tot, g["vol"], f"=SUM({col_vol}{FILA_INICIO}:{col_vol}{ultima_fila})", "0.000", FUENTE_TOTAL)
    _escribir(ws, r_tot, (g["lat_g"][0], ultima_col), None)

    r_nota = r_tot + 1
    _escribir(ws, r_nota, (1, ultima_col), "Imprimir en tamaño 8.5x13", alineacion=IZQUIERDA)
    for c in range(1, ultima_col + 1):
        ws.cell(r_nota, c).border = Border()

    # --- Columna ID: encabezado, angosta y agrupada/oculta ------------------
    for r in (8, 9, 10):
        copiar_estilo(ws.cell(r, COL_DAP_C), ws.cell(r, col_id))
    ws.cell(8, col_id).value = "ID"
    ws.merge_cells(start_row=8, start_column=col_id, end_row=10, end_column=col_id)
    letra_id = L(col_id)
    ws.column_dimensions[letra_id].width = 6
    ws.column_dimensions.group(letra_id, letra_id, hidden=True, outline_level=1)

    # --- Impresión (8.5 x 13 = Folio) ---------------------------------------
    ws.print_area = f"A1:{L(ultima_col)}{r_nota}"
    ws.print_title_rows = "8:10"
    ws.page_setup.paperSize = 14  # Folio 8.5 x 13 in
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.verticalCentered = False

    # Sin líneas internas en celdas combinadas (se ven en Google Sheets / LibreOffice)
    for hoja in wb.worksheets:
        normalizar_bordes_combinadas(hoja)

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()

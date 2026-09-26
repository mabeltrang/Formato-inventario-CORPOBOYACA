"""Utilidades de openpyxl que la librería no resuelve sola.

`insert_cols` de openpyxl mueve valores y estilos, pero NO mueve las celdas combinadas
ni los anchos de columna. Estas funciones lo hacen.
"""

from __future__ import annotations

from copy import copy

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension


def anchos_columnas(ws, hasta: int) -> dict[int, float | None]:
    """Ancho de cada columna 1..hasta, expandiendo los rangos min/max de openpyxl."""
    anchos: dict[int, float | None] = {}
    for dim in ws.column_dimensions.values():
        if dim.min is None:
            continue
        for c in range(dim.min, (dim.max or dim.min) + 1):
            anchos[c] = dim.width
    return {c: anchos.get(c) for c in range(1, hasta + 1)}


def fijar_anchos(ws, anchos: dict[int, float | None]) -> None:
    """Reescribe los anchos columna por columna."""
    ws.column_dimensions.clear()
    for c, w in anchos.items():
        letra = get_column_letter(c)
        dim = ColumnDimension(ws, index=letra)
        if w is not None:
            dim.width = w
        ws.column_dimensions[letra] = dim


def insertar_columnas(ws, antes_de: int, cantidad: int, ancho: float | None = None) -> None:
    """Inserta `cantidad` columnas antes de la columna `antes_de` (1-based).

    Las combinaciones que empiezan en o después de `antes_de` se desplazan;
    las que la atraviesan se ensanchan.
    """
    if cantidad <= 0:
        return
    max_col = max(ws.max_column, 60)
    anchos = anchos_columnas(ws, max_col)

    rangos = [(m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges]
    for m in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(m))

    ws.insert_cols(antes_de, cantidad)

    for r1, c1, r2, c2 in rangos:
        if c1 >= antes_de:
            c1, c2 = c1 + cantidad, c2 + cantidad
        elif c2 >= antes_de:
            c2 += cantidad
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

    nuevos = {}
    for c, w in anchos.items():
        nuevos[c + cantidad if c >= antes_de else c] = w
    for i in range(cantidad):
        nuevos[antes_de + i] = ancho if ancho is not None else anchos.get(antes_de - 1)
    fijar_anchos(ws, nuevos)


def copiar_estilo(origen, destino) -> None:
    destino.font = copy(origen.font)
    destino.border = copy(origen.border)
    destino.fill = copy(origen.fill)
    destino.alignment = copy(origen.alignment)
    destino.number_format = origen.number_format


def normalizar_bordes_combinadas(ws) -> None:
    """Deja los bordes solo en el contorno de cada rango combinado.

    Excel ignora los bordes internos de una celda combinada, pero Google Sheets y
    LibreOffice los dibujan (líneas verticales atravesando el texto). Aquí se toma el
    borde de la celda ancla y se reparte: izquierdo en la primera columna, derecho en
    la última, superior en la primera fila e inferior en la última.
    """
    from openpyxl.styles import Border, Side

    vacio = Side()
    for rango in ws.merged_cells.ranges:
        ancla = ws.cell(rango.min_row, rango.min_col)
        b = ancla.border
        izquierdo, superior = b.left or vacio, b.top or vacio
        # El borde derecho/inferior puede estar guardado en la última celda y no en el ancla
        ub = ws.cell(rango.max_row, rango.max_col).border
        derecho = b.right if (b.right and b.right.style) else (ub.right or vacio)
        inferior = b.bottom if (b.bottom and b.bottom.style) else (ub.bottom or vacio)
        for r in range(rango.min_row, rango.max_row + 1):
            for c in range(rango.min_col, rango.max_col + 1):
                ws.cell(r, c).border = Border(
                    left=izquierdo if c == rango.min_col else vacio,
                    right=derecho if c == rango.max_col else vacio,
                    top=superior if r == rango.min_row else vacio,
                    bottom=inferior if r == rango.max_row else vacio,
                )

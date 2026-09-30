"""Lectura del inventario forestal en el formato propio (hoja "Inventario").

Estructura esperada:
    - Encabezado del proyecto en las primeras filas (PROYECTO, PROPIETARIO, UBICACIÓN, FECHA).
    - Una fila de encabezados con: ID, Nombre común, Nombre científico, CAP A (m), CAP B (m)...,
      HT (m), VT (m3), Coordenada X (longitud decimal), Coordenada Y (latitud decimal).
    - Una fila por árbol. Los fustes adicionales van en columnas CAP B, CAP C, CAP D...
"""

from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass, field

import openpyxl

from .calculos import FACTOR_FORMA_DEFECTO, area_basal, dap_desde_cap, volumen_total

# Tolerancia para comparar el VT recalculado con el del Excel de origen (m³)
TOLERANCIA_VT = 0.0005

# Caja aproximada de Colombia continental, para detectar coordenadas absurdas
LAT_MIN, LAT_MAX = -4.5, 13.5
LON_MIN, LON_MAX = -79.5, -66.5


class ErrorFormato(ValueError):
    """El archivo no tiene la estructura esperada."""


@dataclass
class Arbol:
    fila: int
    id: object
    nombre_comun: str
    nombre_cientifico: str
    familia: str
    caps: list[float]          # CAP de cada fuste (m), en orden A, B, C...
    altura: float              # HT (m)
    lon: float                 # Coordenada X (grados decimales)
    lat: float                 # Coordenada Y (grados decimales)
    vt_origen: float | None    # VT que traía el Excel de origen (si tenía valor calculado)
    vt: float = 0.0            # VT recalculado

    @property
    def daps_cm(self) -> list[float]:
        return [dap_desde_cap(c) * 100 for c in self.caps]

    @property
    def abt(self) -> float:
        return sum(area_basal(c) for c in self.caps)


@dataclass
class Inventario:
    proyecto: str = ""
    propietario: str = ""
    ubicacion: str = ""
    fecha: str = ""
    hoja: str = ""
    arboles: list[Arbol] = field(default_factory=list)
    errores: list[str] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)

    @property
    def max_fustes(self) -> int:
        return max((len(a.caps) for a in self.arboles), default=1)


# ----------------------------------------------------------------------------
# Utilidades
# ----------------------------------------------------------------------------

def _norm(texto) -> str:
    """Minúsculas, sin tildes y con espacios simples."""
    if texto is None:
        return ""
    t = unicodedata.normalize("NFKD", str(texto))
    t = "".join(ch for ch in t if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", t).strip().lower()


def _num(valor):
    """Convierte a float aceptando coma decimal. Devuelve None si no es numérico."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    t = str(valor).strip().replace(" ", "")
    if not t or t.startswith("="):
        return None
    t = t.replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


RE_CAP = re.compile(r"^cap\s*([a-z])\b")

ALIAS = {
    "id": ["id"],
    "nombre_comun": ["nombre comun"],
    "nombre_cientifico": ["nombre cientifico"],
    "familia": ["familia"],
    "ht": ["ht (m)", "ht", "altura total (m)", "altura total"],
    "vt": ["vt (m3)", "vt (m³)", "vt"],
    "lon": ["coordenada x", "coord x", "longitud"],
    "lat": ["coordenada y", "coord y", "latitud"],
}


def _buscar_encabezados(ws) -> tuple[int, dict, list[tuple[str, int]]]:
    """Encuentra la fila de encabezados y devuelve (fila, columnas, columnas_CAP)."""
    for r in range(1, min(ws.max_row, 30) + 1):
        textos = {c: _norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)}
        if "id" not in textos.values() or not any(RE_CAP.match(t) for t in textos.values()):
            continue
        cols = {}
        for clave, alias in ALIAS.items():
            for c, t in textos.items():
                if t in alias:
                    cols[clave] = c
                    break
        caps = sorted(
            ((RE_CAP.match(t).group(1).upper(), c) for c, t in textos.items() if RE_CAP.match(t)),
            key=lambda x: x[0],
        )
        # Unidad del CAP según el encabezado de "CAP A": (cm) o (m). Sin unidad → se decide por los datos.
        texto_cap_a = next((t for t in textos.values() if RE_CAP.match(t) and RE_CAP.match(t).group(1) == "a"), "")
        if "cm" in texto_cap_a:
            cols["_unidad_cap"] = "cm"
        elif "(m)" in texto_cap_a:
            cols["_unidad_cap"] = "m"
        return r, cols, caps
    raise ErrorFormato(
        "No encontré la fila de encabezados (debe tener 'ID' y al menos una columna 'CAP A (m)')."
    )


def _metadatos(ws, fila_enc: int) -> dict:
    etiquetas = {"proyecto": "proyecto", "propietario": "propietario", "ubicacion": "ubicacion", "fecha": "fecha"}
    meta = {}
    for r in range(1, fila_enc):
        for c in range(1, ws.max_column):
            t = _norm(ws.cell(r, c).value)
            if t in etiquetas:
                valor = ws.cell(r, c + 1).value
                meta[etiquetas[t]] = "" if valor is None else str(valor).strip()
    return meta


# ----------------------------------------------------------------------------
# Lectura principal
# ----------------------------------------------------------------------------

def leer_inventario(archivo, factor_forma: float = FACTOR_FORMA_DEFECTO) -> Inventario:
    """Lee el inventario. `archivo` puede ser una ruta o un objeto tipo archivo (bytes)."""
    datos = archivo.read() if hasattr(archivo, "read") else open(archivo, "rb").read()
    wb_f = openpyxl.load_workbook(io.BytesIO(datos))                   # fórmulas
    wb_v = openpyxl.load_workbook(io.BytesIO(datos), data_only=True)   # valores calculados

    nombre_hoja = next((n for n in wb_f.sheetnames if _norm(n) == "inventario"), wb_f.sheetnames[0])
    ws_f, ws_v = wb_f[nombre_hoja], wb_v[nombre_hoja]

    fila_enc, cols, caps_cols = _buscar_encabezados(ws_f)
    faltan = [k for k in ("nombre_comun", "nombre_cientifico", "ht", "lon", "lat") if k not in cols]
    if faltan:
        raise ErrorFormato(f"Faltan columnas en la hoja '{nombre_hoja}': {', '.join(faltan)}")
    if not caps_cols or caps_cols[0][0] != "A":
        raise ErrorFormato("Falta la columna 'CAP A (m)'.")

    inv = Inventario(hoja=nombre_hoja, **_metadatos(ws_v, fila_enc))

    # Factor para llevar el CAP a metros
    unidad_cap = cols.pop("_unidad_cap", None)
    if unidad_cap is None:
        muestra = [
            _num(ws_v.cell(r, caps_cols[0][1]).value) or _num(ws_f.cell(r, caps_cols[0][1]).value)
            for r in range(fila_enc + 1, min(ws_f.max_row, fila_enc + 60) + 1)
        ]
        muestra = sorted(v for v in muestra if v and v > 0)
        unidad_cap = "cm" if muestra and muestra[len(muestra) // 2] > 5 else "m"
        inv.avisos.append(
            f"El encabezado de CAP A no indica unidad; por los valores se asumió CAP en {unidad_cap}."
        )
    factor_cap = 0.01 if unidad_cap == "cm" else 1.0

    def valor(r, clave):
        c = cols.get(clave)
        if c is None:
            return None
        v = ws_v.cell(r, c).value
        return v if v is not None else ws_f.cell(r, c).value

    id_anterior = None
    for r in range(fila_enc + 1, ws_f.max_row + 1):
        nombre = valor(r, "nombre_comun")
        cap_a = _num(ws_v.cell(r, caps_cols[0][1]).value) or _num(ws_f.cell(r, caps_cols[0][1]).value)
        if (nombre is None or str(nombre).strip() == "") and cap_a is None:
            continue  # fila vacía o fila de totales

        # ID: valor calculado; si es una fórmula sin valor en caché, se infiere del anterior
        id_v = ws_v.cell(r, cols["id"]).value if "id" in cols else None
        if id_v is None and "id" in cols:
            crudo = ws_f.cell(r, cols["id"]).value
            if isinstance(crudo, str) and crudo.startswith("=") and isinstance(id_anterior, (int, float)):
                id_v = int(id_anterior) + 1
                inv.avisos.append(f"Fila {r}: el ID es una fórmula sin valor guardado; se asumió {id_v}.")
            else:
                id_v = crudo
        if isinstance(id_v, float) and id_v.is_integer():
            id_v = int(id_v)
        id_anterior = id_v
        ref = f"ID {id_v} (fila {r})"

        # Fustes
        caps = []
        for letra, c in caps_cols:
            v = _num(ws_v.cell(r, c).value)
            if v is None:
                v = _num(ws_f.cell(r, c).value)
            if v is not None and v > 0:
                caps.append(v * factor_cap)
            elif v is not None and v < 0:
                inv.errores.append(f"{ref}: CAP {letra} negativo ({v}).")
        if not caps:
            inv.errores.append(f"{ref}: no tiene CAP A.")
            continue

        altura = _num(valor(r, "ht"))
        if altura is None or altura <= 0:
            inv.errores.append(f"{ref}: altura total (HT) vacía o no válida.")
            continue

        lon, lat = _num(valor(r, "lon")), _num(valor(r, "lat"))
        if lon is None or lat is None:
            inv.errores.append(f"{ref}: coordenadas vacías.")
            continue
        # Coordenadas invertidas (X con la latitud y Y con la longitud)
        if LAT_MIN <= lon <= LAT_MAX and LON_MIN <= lat <= LON_MAX:
            lon, lat = lat, lon
            inv.avisos.append(f"{ref}: Coordenada X y Y parecían invertidas; se intercambiaron.")
        if not (LAT_MIN <= lat <= LAT_MAX and LON_MIN <= lon <= LON_MAX):
            inv.errores.append(f"{ref}: coordenadas fuera de Colombia (X={lon}, Y={lat}).")
            continue

        vt_origen = _num(ws_v.cell(r, cols["vt"]).value) if "vt" in cols else None

        arbol = Arbol(
            fila=r,
            id=id_v,
            nombre_comun=str(nombre or "").strip(),
            nombre_cientifico=str(valor(r, "nombre_cientifico") or "").strip(),
            familia=str(valor(r, "familia") or "").strip(),
            caps=caps,
            altura=altura,
            lon=lon,
            lat=lat,
            vt_origen=vt_origen,
        )
        arbol.vt = volumen_total(caps, altura, factor_forma)
        if vt_origen is not None and abs(arbol.vt - vt_origen) > TOLERANCIA_VT:
            inv.avisos.append(
                f"{ref}: el VT del Excel ({vt_origen:.4f} m³) no coincide con el recalculado "
                f"({arbol.vt:.4f} m³). Se usa el recalculado."
            )
        inv.arboles.append(arbol)

    if not inv.arboles:
        inv.errores.append("No se encontró ningún árbol válido en el inventario.")

    # IDs repetidos
    vistos = {}
    for a in inv.arboles:
        if a.id is not None and a.id in vistos:
            inv.avisos.append(f"ID {a.id} repetido (filas {vistos[a.id]} y {a.fila}).")
        vistos.setdefault(a.id, a.fila)

    return inv

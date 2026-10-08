import io

import openpyxl
import pytest

from fgr06.costos import (DatosFGR29, Tarifas, costo_aprovechamiento, costo_compensacion, generar_fgr29,
                          generar_tablas_informe, totales_fgr29)


def test_aprovechamiento_por_volumen():
    items = costo_aprovechamiento(6.508, Tarifas())
    assert sum(i.total for i in items) == pytest.approx(6.508 * 190_000)


def test_compensacion_igual_al_formato_estandar():
    # Formato "Costos compensación" con 190 plantas
    c = costo_compensacion(190, Tarifas())
    assert c.subtotales == {
        "Mano de obra siembra inicial": 3_515_000,
        "Insumos": 10_099_000,
        "Herramientas": 742_000,
        "Reposición, mantenimiento y monitoreo (3 años)": 10_413_095,
    }
    assert c.imprevistos == pytest.approx(520_654.75)
    assert c.total == pytest.approx(25_289_749.75)


def test_compensacion_escala_con_plantas():
    t = Tarifas()
    assert costo_compensacion(470, t).total > costo_compensacion(141, t).total


def test_fgr29_escritura():
    t = Tarifas()
    comp = costo_compensacion(470, t)
    d = DatosFGR29(6.508, comp, t, valor_predio=84_000_000, canon_anual=6_800_000)
    wb = openpyxl.load_workbook(io.BytesIO(generar_fgr29(d)))
    a, b = wb.worksheets
    assert a["D10"].value == 6.508 and a["E10"].value == 190_000
    assert a["E42"].value == 0                    # la compensación solo va en 2.6
    assert a["E39"].value == 84_000_000 and b["E28"].value == 6_800_000
    seccion = sum(b[c].value for c in ("E51", "E52", "E53", "E54", "E55"))
    assert seccion == pytest.approx(comp.total, abs=5)


def test_tablas_informe():
    t = Tarifas()
    wb = openpyxl.load_workbook(io.BytesIO(generar_tablas_informe(6.508, costo_compensacion(141, t), t)))
    assert wb.sheetnames == ["Costos aprovechamiento", "Costos reposición"]


def test_totales_fgr29_iguales_a_formulas():
    t = Tarifas()
    d = DatosFGR29(6.508, costo_compensacion(470, t), t, valor_predio=84_000_000, canon_anual=6_800_000)
    tot = totales_fgr29(d)
    assert tot["inversion"] == pytest.approx(289_476_520)
    assert tot["total"] == tot["inversion"] + tot["operacion"]

import io
import math
from pathlib import Path

import openpyxl
import pytest

from fgr06 import generar_fgr06, leer_inventario
from fgr06.calculos import area_basal, decimal_a_gms, volumen_total

EJEMPLO = Path(__file__).resolve().parent.parent / "ejemplos" / "inventario_ejemplo.xlsx"


def test_area_basal_y_volumen():
    # CAP 1.02 m -> DAP 0.32468 m -> AB 0.082792 m²
    assert area_basal(1.02) == pytest.approx(0.0827924, rel=1e-6)
    assert volumen_total([1.02], 6) == pytest.approx(0.3477281, rel=1e-6)
    # Varios fustes: misma altura, suma de AB
    assert volumen_total([0.48, 0.45, 0.33], 5) == pytest.approx(0.1509028, rel=1e-6)


def test_gms_con_acarreo():
    assert decimal_a_gms(5.687430555) == (5, 41, 14.75)
    assert decimal_a_gms(-73.04351388) == (73, 2, 36.65)
    assert decimal_a_gms(5 + 41 / 60 + 59.999 / 3600) == (5, 42, 0.0)


def test_lectura_ejemplo():
    inv = leer_inventario(EJEMPLO)
    assert len(inv.arboles) == 5
    assert inv.max_fustes == 5
    assert not inv.errores


def test_fgr06_con_fustes_extra():
    inv = leer_inventario(EJEMPLO)
    wb = openpyxl.load_workbook(io.BytesIO(generar_fgr06(inv)))
    ws = wb.worksheets[0]
    # Con 5 fustes se insertan 2 columnas: DAP D (Q) y DAP E (R)
    assert ws["Q8"].value == "DAP D (cm)"
    assert ws["R8"].value == "DAP E (cm)"
    assert ws["S8"].value == "ALTO TOTAL"
    # Árbol 4 (fila 14) con 5 fustes
    assert ws["U14"].value == pytest.approx(volumen_total([0.52, 0.41, 0.38, 0.30, 0.25], 7))
    assert ws["R14"].value == pytest.approx(0.25 / math.pi * 100)
    # Totales
    assert ws["A16"].value == "=SUM(A11:A15)"
    assert ws["U16"].value == "=SUM(U11:U15)"
    # ID oculto
    assert ws["AP8"].value == "ID"
    assert ws.column_dimensions["AP"].hidden

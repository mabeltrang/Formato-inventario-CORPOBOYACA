import io
import zipfile
from pathlib import Path

import openpyxl
import pytest

from fgr06 import generar_fgr06, leer_inventario
from fgr06.amenazas import consultar_especie
from fgr06.parte_b import arbol_central, calcular_parte_b, jornales
from fgr06.precios import buscar_precio, cargar_precios
from fgr06.predio import _metros_por_grado, calcular_relieve, leer_kmz

EJEMPLO = Path(__file__).resolve().parent.parent / "ejemplos" / "inventario_ejemplo.xlsx"

KML = """<?xml version="1.0"?><kml xmlns="http://www.opengis.net/kml/2.2"><Document><Placemark>
<name>Predio</name><Polygon><outerBoundaryIs><LinearRing><coordinates>
-73.0460,5.6860,0 -73.0425,5.6860,0 -73.0425,5.6885,0 -73.0460,5.6885,0 -73.0460,5.6860,0
</coordinates></LinearRing></outerBoundaryIs></Polygon></Placemark></Document></kml>"""


def _kmz() -> bytes:
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        z.writestr("doc.kml", KML)
    return b.getvalue()


def test_amenazas():
    assert not consultar_especie("Eucalyptus camaldulensis").es_amenazada()
    assert consultar_especie("Quercus humboldtii").es_amenazada()          # MADS VU
    assert consultar_especie("Cedrela montana").es_amenazada()             # CITES II / UICN VU
    assert consultar_especie("Juglans neotropica").veda == "nacional"


def test_precios_respaldo_por_genero():
    df = cargar_precios()
    assert buscar_precio(df, "Alnus acuminata") == (650000.0, "especie")
    assert buscar_precio(df, "Eucalyptus globulus")[1] == "género"
    assert buscar_precio(df, "Fraxinus uhdei") == (None, "sin precio")


def test_kmz_area_y_pendiente():
    predio = leer_kmz(_kmz())
    assert predio.area_ha == pytest.approx(10.72, abs=0.02)
    assert predio.contiene(-73.0435, 5.6874)
    mx, _ = _metros_por_grado(5.687)
    plano = lambda pts: [2900 + 0.10 * (lon + 73.046) * mx for lon, lat in pts]  # 10 %
    rel = calcular_relieve(predio, fuentes=[("plano", plano)])
    assert rel.pendiente_pct == pytest.approx(10.0, abs=0.1)
    assert rel.topografia == "Ondulada"


def test_jornales():
    assert jornales(55, 6.7, 30, 8) == 2
    assert jornales(5, 0.5, 30, 8) == 1


def test_parte_b_escritura():
    inv = leer_inventario(EJEMPLO)
    d = calcular_parte_b(inv, cargar_precios(), area_predio_ha=2.169)
    assert d.n_plantas == 10 * 5          # ninguna especie amenazada en el ejemplo
    assert arbol_central(inv) is not None
    wb = openpyxl.load_workbook(io.BytesIO(generar_fgr06(inv, parte_b=d)))
    ws = wb.worksheets[1]
    assert ws["AC27"].value == 2.169
    assert ws["A34"].value == 3                   # 3 eucaliptos
    assert ws["Q34"].value == "Eucalyptus camaldulensis"
    assert ws["E99"].value == 50
    assert ws["AJ97"].value == "x"

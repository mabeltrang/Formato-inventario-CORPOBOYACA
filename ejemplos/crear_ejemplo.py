"""Genera ejemplos/inventario_ejemplo.xlsx: inventario ficticio en el formato propio,
con un árbol de 5 fustes para probar la inserción de columnas DAP D y DAP E."""

from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter as L

SALIDA = Path(__file__).with_name("inventario_ejemplo.xlsx")

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Inventario"
for i, (k, v) in enumerate([("PROYECTO", "Proyecto de ejemplo"), ("PROPIETARIO", "Propietario ficticio"),
                            ("UBICACIÓN", "Firavitoba, Boyacá"), ("FECHA", "Septiembre, 2026")], start=1):
    ws[f"D{i}"], ws[f"E{i}"] = k, v

enc = ["ID", "Cobertura", "Nombre común", "Nombre científico", "Familia"]
for f in "ABCDE":
    enc += [f"CAP {f} (m)", f"DAP {f} (m)", f"AB {f} (m2)"]
enc += ["HT (m)", "HC (m)", "AB T (m2)", "VT (m3)", "VC (m3)", "Coordenada X ", "Coordenada Y",
        "Municipio", "Departamento", "Estado fitosanitario", "Medida de manejo"]
for c, t in enumerate(enc, start=1):
    ws.cell(5, c, t)
col = {t.strip(): i for i, t in enumerate(enc, start=1)}

arboles = [
    (1, "Eucalipto", "Eucalyptus camaldulensis", "Myrtaceae", [1.02], 6, -73.043514, 5.687431),
    (2, "Aliso", "Alnus acuminata", "Betulaceae", [0.47], 3, -73.043403, 5.687422),
    (3, "Eucalipto", "Eucalyptus camaldulensis", "Myrtaceae", [0.48, 0.45, 0.33], 5, -73.043578, 5.686481),
    (4, "Eucalipto", "Eucalyptus camaldulensis", "Myrtaceae", [0.52, 0.41, 0.38, 0.30, 0.25], 7, -73.044100, 5.686900),
    (5, "Pino", "Pinus sp", "Pinaceae", [0.95], 3, -73.043550, 5.686367),
]
for i, (id_, comun, cient, fam, caps, ht, x, y) in enumerate(arboles):
    r = 6 + i
    ws.cell(r, col["ID"], id_)
    ws.cell(r, col["Cobertura"], "Pastos limpios")
    ws.cell(r, col["Nombre común"], comun)
    ws.cell(r, col["Nombre científico"], cient)
    ws.cell(r, col["Familia"], fam)
    abs_ = []
    for f, cap in zip("ABCDE", caps):
        c_cap = col[f"CAP {f} (m)"]
        ws.cell(r, c_cap, cap)
        ws.cell(r, c_cap + 1, f"={L(c_cap)}{r}/PI()")
        ws.cell(r, c_cap + 2, f"=PI()*({L(c_cap + 1)}{r}/2)^2")
        abs_.append(f"{L(c_cap + 2)}{r}")
    ws.cell(r, col["HT (m)"], ht)
    ws.cell(r, col["AB T (m2)"], "=" + "+".join(abs_))
    ws.cell(r, col["VT (m3)"], f"={L(col['HT (m)'])}{r}*{L(col['AB T (m2)'])}{r}*0.7")
    ws.cell(r, col["Coordenada X"], x)
    ws.cell(r, col["Coordenada Y"], y)
    ws.cell(r, col["Municipio"], "Firavitoba")
    ws.cell(r, col["Departamento"], "Boyacá")
    ws.cell(r, col["Estado fitosanitario"], "Bueno")
    ws.cell(r, col["Medida de manejo"], "Tala")

wb.save(SALIDA)
print(f"Creado {SALIDA}")
